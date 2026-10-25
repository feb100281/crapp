"""
Касса: карточки счёта 50 из 1С (Excel) → строки для parquet.

Полная пересборка при каждом запуске: все .xlsx папки (подпапки допускаются),
строки как есть, без ключей и слияний. Файлы должны покрывать закрытые
периоды без пересечений (например, по месяцу на файл). Пересечение периодов
и разрыв сальдо между соседними файлами не останавливают загрузку — пишутся
в лог как предупреждения, чтобы косяки кассы были видны.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

DOC = re.compile(r"^(?P<type>.*?)\s+(?P<num>\S+)\s+от\s+(?P<d>\d{2}\.\d{2}\.\d{4})(?:\s+(?P<t>\d{2}:\d{2}:\d{2}))?")
TITLE = re.compile(r"Карточка счета\s+(?P<acc>[\d.]+)\s+за\s+(?P<period>.+?)\s*$")
DIVISION = re.compile(r"\bОП\s+[\"«]([^\"»]+)[\"»]")
DAY = re.compile(r"^\d{2}\.\d{2}\.\d{4}$")

COLUMNS = {                       # колонки parquet строк
    "source_file": "VARCHAR",
    "cash_account": "VARCHAR", "date": "DATE",
    "doc_type": "VARCHAR", "doc_number": "VARCHAR", "doc_date": "DATE", "doc_datetime": "TIMESTAMP",
    "doc_comment": "VARCHAR",
    "division": "VARCHAR",
    "direction": "INTEGER",       # 1 — приход в кассу, 2 — расход из кассы
    "account_dt": "VARCHAR", "account_kt": "VARCHAR", "corr_account": "VARCHAR",
    "amount": "DECIMAL(18,2)",
    "cf_item_1c": "VARCHAR",      # аналитика на стороне 50 — статья движения ДС в 1С
    "corr_subconto1": "VARCHAR", "corr_subconto2": "VARCHAR", "corr_subconto3": "VARCHAR",
    "analytics_dt": "VARCHAR", "analytics_kt": "VARCHAR",
    "balance": "DECIMAL(18,2)",   # текущее сальдо после строки, дебетовое со знаком +
}
FILE_COLUMNS = {                  # итоги файлов — для сверки
    "source_file": "VARCHAR", "cash_account": "VARCHAR", "title": "VARCHAR",
    "date_from": "DATE", "date_to": "DATE", "rows": "INTEGER",
    "opening": "DECIMAL(18,2)", "debit": "DECIMAL(18,2)", "credit": "DECIMAL(18,2)", "closing": "DECIMAL(18,2)",
}


def _num(value) -> float | None:
    if value in (None, ""):
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def _signed(side, value) -> float | None:
    v = _num(value)
    return None if v is None else (-v if str(side).strip().upper().startswith("К") else v)


def _lines(text) -> list[str]:
    return [s.strip() for s in str(text or "").split("\n")]


def _header(ws) -> tuple[int, dict[str, int], str]:
    title = ""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if not title and cells and TITLE.match(cells[0]):
            title = cells[0]
        if "Период" in cells and "Документ" in cells:
            cols = {name: cells.index(name) for name in
                    ("Период", "Документ", "Аналитика Дт", "Аналитика Кт", "Дебет", "Кредит", "Текущее сальдо")}
            return i, cols, title
    raise ValueError("не найдена шапка карточки счёта (Период / Документ / Дебет / Кредит)")


def parse_file(path: Path, root: Path) -> tuple[list[dict], dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    head_row, c, title = _header(ws)
    m = TITLE.match(title)
    cash_account = m.group("acc") if m else "50"
    rel = path.relative_to(root).as_posix()
    rows, info = [], {"source_file": rel, "cash_account": cash_account, "title": title,
                      "opening": None, "debit": None, "credit": None, "closing": None}

    for r in ws.iter_rows(min_row=head_row + 1, values_only=True):
        first = str(r[c["Период"]] or "").strip()
        if first.startswith("Сальдо на начало"):
            info["opening"] = _signed(r[c["Текущее сальдо"]], r[c["Текущее сальдо"] + 1])
            continue
        if first.startswith("Обороты"):
            info["debit"], info["credit"] = _num(r[c["Дебет"]]), _num(r[c["Кредит"]])
            info["closing"] = _signed(r[c["Текущее сальдо"]], r[c["Текущее сальдо"] + 1])
            continue
        if not DAY.match(first):
            continue

        acc_dt, acc_kt = str(r[c["Дебет"]] or "").strip(), str(r[c["Кредит"]] or "").strip()
        incoming = acc_dt.startswith("50")
        amount = _num(r[c["Дебет"] + 1]) if incoming else _num(r[c["Кредит"] + 1])
        doc = _lines(r[c["Документ"]])
        dm = DOC.match(doc[0]) if doc else None
        an_dt, an_kt = str(r[c["Аналитика Дт"]] or "").strip(), str(r[c["Аналитика Кт"]] or "").strip()
        own, corr = (an_dt, an_kt) if incoming else (an_kt, an_dt)
        sub = (_lines(corr) + ["", "", ""])[:3]
        comment = " ".join(s for s in doc[1:] if s).strip(" ,")
        rows.append({
            "source_file": rel, "cash_account": cash_account,
            "date": datetime.strptime(first, "%d.%m.%Y").date(),
            "doc_type": dm.group("type") if dm else (doc[0] if doc else ""),
            "doc_number": dm.group("num") if dm else "",
            "doc_date": datetime.strptime(dm.group("d"), "%d.%m.%Y").date() if dm else None,
            "doc_datetime": datetime.strptime(f"{dm.group('d')} {dm.group('t') or '00:00:00'}", "%d.%m.%Y %H:%M:%S")
            if dm else None,
            "doc_comment": comment,
            "division": (DIVISION.search(comment).group(1) if DIVISION.search(comment) else ""),
            "direction": 1 if incoming else 2,
            "account_dt": acc_dt, "account_kt": acc_kt, "corr_account": acc_kt if incoming else acc_dt,
            "amount": amount,
            "cf_item_1c": own, "corr_subconto1": sub[0], "corr_subconto2": sub[1], "corr_subconto3": sub[2],
            "analytics_dt": an_dt, "analytics_kt": an_kt,
            "balance": _signed(r[c["Текущее сальдо"]], r[c["Текущее сальдо"] + 1]),
        })
    wb.close()


    days = [x["date"] for x in rows]
    info.update(date_from=min(days) if days else None, date_to=max(days) if days else None, rows=len(rows))
    return rows, info


def check_files(files: list[dict]) -> list[str]:
    """Пересечения периодов и разрывы сальдо между соседними файлами одного счёта."""
    problems = []
    by_acc: dict[str, list[dict]] = {}
    for f in files:
        by_acc.setdefault(f["cash_account"], []).append(f)
    for acc, items in by_acc.items():
        items.sort(key=lambda f: (f["date_from"] or date.min, f["date_to"] or date.min))
        for prev, cur in zip(items, items[1:]):
            if cur["date_from"] and prev["date_to"] and cur["date_from"] <= prev["date_to"]:
                problems.append(f"{acc}: периоды пересекаются — {prev['source_file']} "
                                f"({prev['date_from']} — {prev['date_to']}) и {cur['source_file']} "
                                f"({cur['date_from']} — {cur['date_to']})")
            elif None not in (prev["closing"], cur["opening"]) and abs(prev["closing"] - cur["opening"]) >= 0.01:
                problems.append(f"{acc}: разрыв сальдо — конец {prev['source_file']} {prev['closing']}, "
                                f"начало {cur['source_file']} {cur['opening']}")
    return problems


def parse_folder(root: Path, log=print) -> tuple[list[dict], list[dict]]:
    root = Path(root)
    files = sorted(p for p in root.rglob("*.xlsx")
                   if p.is_file() and not p.name.startswith(("~$", ".")))
    log(f"Файлов Excel: {len(files)}")
    out, infos, bad = [], [], []
    for p in files:
        try:
            rows, info = parse_file(p, root)
        except Exception as exc:  # noqa: BLE001
            bad.append(f"{p.relative_to(root)}: {exc}")
            continue
        log(f"  {info['source_file']}: {info['title']} · строк {info['rows']} · "
            f"{info['date_from']} — {info['date_to']}")
        out.extend(rows)
        infos.append(info)
    for problem in bad + check_files(infos):
        log(f"  ВНИМАНИЕ: {problem}")
    out.sort(key=lambda r: (r["date"], r["doc_datetime"] or datetime.min, r["doc_number"]))
    log(f"Строк: {len(out)}")
    return out, infos
