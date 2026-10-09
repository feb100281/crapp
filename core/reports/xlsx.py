"""
Книга Excel для выгрузок: оглавление с KPI и навигацией + листы-таблицы.

    book = Book("Отчёт ДДС", subtitle="…", params="RUB · период …")
    book.kpi("Чистый поток", 1_250_000, "2025 год")
    sheet = book.sheet("ДДС", "Движение денежных средств", description="…")
    sheet.table(
        [Col("Код", kind="code"), Col("Статья", width=44), Col("Итого", kind="money", total=True)],
        [Row(["110100", "Выручка", 100.0], level=1), …],
        total=Row(["", "ИТОГО", 100.0]),
    )
    data = book.save()

Отрицательные числа — в скобках, ноль — тире. Уровни иерархии различаются
заливкой и отступом в колонке с indent=True, вложенные строки группируются
(плюсики слева, Row.outline / hidden / collapsed), колонки — так же
(плюсики сверху, Col.outline / hidden / collapsed; итог группы слева). Зебра не перекрывает
смысловые заливки (итоги, уровни, «Итого»). Коды из цифр пишутся числом —
без зелёных треугольников «число сохранено как текст».
"""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ---------------------------------------------------------------- палитра
NAVY = "2F6656"
NAVY_2 = "3D7A67"
NAVY_3 = "1F5E4E"
TEXT = "1F1F1F"
TEXT_2 = "4A4A4A"
MUTED = "8A8A8A"
SURFACE_2 = "F3F8F6"
SURFACE_3 = "EDF5F1"
SURFACE_4 = "E7F1ED"
TOTAL_ROW = "E7F1ED"
ZEBRA_ROW = "F7F7F7"
LINE = "D9D9D9"
INCOME = "3C4043"
EXPENSE = "7B4437"
WARN = "9A6100"
WARN_BG = "FDF3DE"

FONT = "Roboto"

# ------------------------------------------------------- форматы чисел
FMT_MONEY = '#,##0;(#,##0);"–"'
FMT_MONEY_DEC = '#,##0.00;(#,##0.00);"–"'
FMT_PCT = '#,##0.0" %";(#,##0.0" %");"–"'
FMT_QTY = '#,##0;(#,##0);"–"'
FMT_DATE = "dd.mm.yyyy"
FMT_DATETIME = "dd.mm.yyyy hh:mm"
FMT_CODE = "0"

TOC_SHEET_NAME = "Оглавление"
HEADER_ROW = 6          # строки 1–5 — шапка листа

NUMBER_FORMATS = {
    "money": FMT_MONEY,
    "money_dec": FMT_MONEY_DEC,
    "pct": FMT_PCT,
    "int": FMT_QTY,
    "date": FMT_DATE,
    "datetime": FMT_DATETIME,
    "code": FMT_CODE,
}
NUMERIC = ("money", "money_dec", "pct", "int")

LEVEL_FILL = {1: SURFACE_4, 2: SURFACE_3, 3: SURFACE_2, "warn": WARN_BG}
LEVEL_BOLD = {1: True, 2: True}


@dataclass
class Col:
    title: str
    kind: str = "text"          # text | code | money | money_dec | int | pct | date | datetime
    width: float | None = None
    total: bool = False         # колонка «Итого» — своя заливка
    wrap: bool = False
    indent: bool = False        # сюда идёт отступ уровня строки (Row.indent)
    outline: int = 0            # группировка колонок (плюсики сверху), например дни месяца
    hidden: bool = False        # колонка свёрнута
    collapsed: bool = False     # у колонки свёрнуты вложенные (она — итог группы слева)


@dataclass
class Row:
    values: list
    level: int | str | None = None   # 1–3 — уровень иерархии, "warn" — внимание,
                                     # "band" — тёмная плашка (остатки), "total" — итог
    bold: bool = False
    outline: int = 0                 # уровень группировки Excel (плюсики слева)
    indent: int = 0                  # отступ в колонке Col(indent=True)
    hidden: bool = False             # строка свёрнута (раскрывается плюсиком)
    collapsed: bool = False          # у строки свёрнуты вложенные


@dataclass
class _Kpi:
    label: str
    value: object
    note: str = ""
    fmt: str = FMT_MONEY


def _side(color: str = LINE, style: str = "thin") -> Side:
    return Side(style=style, color=color)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", start_color=color, end_color=color)


def _clean(value):
    """Значение ячейки: Decimal → float, дата со временем — без часового пояса."""
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    return value


class Sheet:
    def __init__(self, book: "Book", ws, title: str, subtitle: str, params: str, description: str):
        self.book = book
        self.ws = ws
        self.title = title
        self.subtitle = subtitle
        self.params = params
        self.description = description
        self.width = 2
        self.outlined = 0

        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_options.horizontalCentered = True

    # ------------------------------------------------------------------

    def _header(self, last_col: int) -> None:
        ws = self.ws
        last_col = max(last_col, 2)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        back = ws.cell(row=1, column=1, value="←  Оглавление")
        back.hyperlink = f"#'{TOC_SHEET_NAME}'!A1"
        back.font = Font(name=FONT, size=9, bold=True, color=NAVY_3, underline=None)
        back.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in (1, 2):
            cell = ws.cell(row=1, column=col)
            cell.fill = _fill(SURFACE_4)
            cell.border = Border(bottom=_side())
        ws.row_dimensions[1].height = 16

        rows = (
            (2, self.title, Font(name=FONT, size=14, bold=True, color=TEXT), 24),
            (3, self.subtitle, Font(name=FONT, size=9, color=TEXT_2), 14),
            (4, self.params, Font(name=FONT, size=9, color=MUTED), 14),
        )
        for row, text, font, height in rows:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row].height = height

        for col in range(1, last_col + 1):
            ws.cell(row=5, column=col).border = Border(bottom=_side(NAVY, "medium"))
        ws.row_dimensions[5].height = 6

    # ------------------------------------------------------------------

    def table(self, cols: list[Col], rows, total: Row | None = None, *,
              freeze_cols: int = 1, freeze_rows: int = 0, autofilter: bool = False,
              note: str = "") -> int:
        """Шапка листа + таблица с шестой строки. Возвращает номер последней строки.
        freeze_rows — сколько первых строк данных закрепить вместе с шапкой."""
        ws, book = self.ws, self.book
        n = len(cols)
        self.width = n
        self._header(n)

        for i, col in enumerate(cols, start=1):
            cell = ws.cell(row=HEADER_ROW, column=i, value=col.title)
            cell._style = book.style(
                bold=True, color="FFFFFF", fill=NAVY, wrap=True, vcenter=True,
                align="right" if col.kind in NUMERIC else "left",
                right=i < n,
            )
        ws.row_dimensions[HEADER_ROW].height = 22

        widths = [len(str(c.title)) for c in cols]
        r = HEADER_ROW
        for r, row in enumerate(rows, start=HEADER_ROW + 1):
            self._write_row(r, cols, row, widths)
        last_data = r

        if total is not None:
            last_data += 1
            self._write_row(last_data, cols, total, widths, is_total=True)

        col_outline = 0
        for i, col in enumerate(cols, start=1):
            width = col.width or min(widths[i - 1] + 2, 38)
            dim = ws.column_dimensions[get_column_letter(i)]
            dim.width = max(width, 6)
            if col.outline:
                dim.outlineLevel = col.outline
                col_outline = max(col_outline, col.outline)
            if col.hidden:
                dim.hidden = True
            if col.collapsed:
                dim.collapsed = True
        if col_outline:
            ws.sheet_properties.outlinePr.summaryRight = False
            ws.sheet_format.outlineLevelCol = col_outline

        if self.outlined:
            ws.sheet_properties.outlinePr.summaryBelow = False
            ws.sheet_format.outlineLevelRow = self.outlined

        ws.freeze_panes = ws.cell(row=HEADER_ROW + 1 + freeze_rows, column=freeze_cols + 1)
        if autofilter and last_data > HEADER_ROW:
            ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(n)}{last_data}"

        if note:
            cell = ws.cell(row=last_data + 2, column=1, value=note)
            cell.font = Font(name=FONT, size=8, color=MUTED)
        return last_data

    def _write_row(self, r: int, cols: list[Col], row: Row, widths: list[int], is_total: bool = False):
        ws, book = self.ws, self.book
        n = len(cols)
        band = row.level == "band"
        is_total = is_total or row.level == "total"
        level_fill = NAVY if band else (TOTAL_ROW if is_total else LEVEL_FILL.get(row.level))
        bold = band or is_total or row.bold or LEVEL_BOLD.get(row.level, False)
        zebra = r % 2 == 0

        for i, col in enumerate(cols, start=1):
            value = _clean(row.values[i - 1]) if i <= len(row.values) else None
            if col.kind == "code" and isinstance(value, str) and value.isdigit() and len(value) <= 15:
                value = int(value)
            numeric = col.kind in NUMERIC
            negative = numeric and isinstance(value, (int, float)) and value < 0

            # смысловая заливка важнее зебры
            fill = level_fill or (SURFACE_3 if col.total else (ZEBRA_ROW if zebra else None))
            if band:
                color = "FFFFFF"
            elif row.level == "warn" and not is_total:
                color = WARN if not numeric else (EXPENSE if negative else TEXT)
            elif col.kind == "code":
                color = MUTED
            elif numeric:
                color = EXPENSE if negative else (TEXT if bold else INCOME)
            else:
                color = TEXT

            cell = ws.cell(row=r, column=i, value=value)
            cell._style = book.style(
                bold=bold or (col.total and numeric), color=color, fill=fill,
                fmt=NUMBER_FORMATS.get(col.kind),
                align="right" if numeric else "left",
                wrap=col.wrap, right=i < n, total=is_total,
                indent=row.indent if col.indent else 0, vcenter=band,
            )
            if value is not None and not col.width:
                size = len(f"{value:,.0f}") if isinstance(value, float) else len(str(value))
                if size > widths[i - 1]:
                    widths[i - 1] = size

        dim = ws.row_dimensions[r]
        if row.outline:
            dim.outlineLevel = row.outline
            self.outlined = max(self.outlined, row.outline)
        if row.hidden:
            dim.hidden = True
        if row.collapsed:
            dim.collapsed = True
        if band:
            dim.height = 20


class Book:
    def __init__(self, title: str, subtitle: str = "", params: str = ""):
        self.title = title
        self.subtitle = subtitle
        self.params = params
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.sheets: list[Sheet] = []
        self.kpis: list[_Kpi] = []
        self._scratch = self.wb.create_sheet("_styles")
        self._styles: dict[tuple, object] = {}

    # --- стили: один раз собираем, дальше раздаём готовые ---

    def style(self, *, bold=False, color=TEXT, fill=None, fmt=None, align="left",
              wrap=False, vcenter=False, right=False, total=False, size=10, indent=0):
        key = (bold, color, fill, fmt, align, wrap, vcenter, right, total, size, indent)
        found = self._styles.get(key)
        if found is None:
            cell = self._scratch.cell(row=len(self._styles) + 1, column=1)
            cell.font = Font(name=FONT, size=size, bold=bold, color=color)
            if fill:
                cell.fill = _fill(fill)
            if fmt:
                cell.number_format = fmt
            cell.alignment = Alignment(
                horizontal=align, vertical="center" if vcenter else "top", wrap_text=wrap,
                indent=indent,
            )
            cell.border = Border(
                right=_side() if right else None,
                top=_side(NAVY, "medium") if total else None,
                bottom=_side(NAVY, "double") if total else None,
            )
            found = self._styles[key] = copy(cell._style)
        return found

    # ------------------------------------------------------------------

    def kpi(self, label: str, value, note: str = "", fmt: str = FMT_MONEY) -> None:
        if len(self.kpis) < 4:
            self.kpis.append(_Kpi(label, value, note, fmt))

    def sheet(self, name: str, title: str | None = None, subtitle: str = "",
              params: str | None = None, description: str = "") -> Sheet:
        safe = "".join(ch for ch in name if ch not in '[]:*?/\\')[:31]
        ws = self.wb.create_sheet(safe)
        sheet = Sheet(self, ws, title or name, subtitle,
                      self.params if params is None else params, description)
        self.sheets.append(sheet)
        return sheet

    # ------------------------------------------------------------------

    def _toc(self) -> None:
        ws = self.wb.create_sheet(TOC_SHEET_NAME, 0)
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.column_dimensions["A"].width = 2
        for letter in "BCDE":
            ws.column_dimensions[letter].width = 26

        head = (
            (1, self.title.upper(), Font(name=FONT, size=14, bold=True, color=TEXT), 24),
            (2, self.subtitle, Font(name=FONT, size=9, color=TEXT_2), 14),
            (3, self.params, Font(name=FONT, size=9, color=MUTED), 14),
        )
        for row, text, font, height in head:
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            cell = ws.cell(row=row, column=2, value=text)
            cell.font = font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row].height = height
        for col in range(2, 6):
            ws.cell(row=4, column=col).border = Border(bottom=_side(NAVY, "medium"))
        ws.row_dimensions[4].height = 6

        r = 6
        if self.kpis:
            for i, k in enumerate(self.kpis):
                col, main = 2 + i, i == 0
                parts = (
                    (r, k.label.upper(), Font(name=FONT, size=8, bold=main, color=NAVY_3 if main else MUTED), None,
                     Border(top=_side(NAVY, "medium"), left=_side(), right=_side())),
                    (r + 1, _clean(k.value), Font(name=FONT, size=16, bold=True, color=NAVY_3), k.fmt,
                     Border(left=_side(), right=_side())),
                    (r + 2, k.note, Font(name=FONT, size=8, color=MUTED), None,
                     Border(bottom=_side(), left=_side(), right=_side())),
                )
                for row, value, font, fmt, border in parts:
                    cell = ws.cell(row=row, column=col, value=value)
                    cell.font = font
                    cell.border = border
                    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
                    if fmt and not isinstance(value, str):
                        cell.number_format = fmt
                    if main:
                        cell.fill = _fill(SURFACE_4)
            for row, height in ((r, 16), (r + 1, 26), (r + 2, 16)):
                ws.row_dimensions[row].height = height
            r += 4

        cell = ws.cell(row=r, column=2, value="СОДЕРЖАНИЕ ОТЧЁТА")
        cell.font = Font(name=FONT, size=10, bold=True, color=TEXT)
        hint = ws.cell(row=r + 1, column=2,
                       value="Щёлкните на названии листа, чтобы перейти. "
                             "На каждом листе есть кнопка возврата в оглавление.")
        hint.font = Font(name=FONT, size=8, color=MUTED)
        r += 2

        for i, sheet in enumerate(self.sheets):
            link = ws.cell(row=r, column=2, value="›  " + sheet.ws.title)
            link.hyperlink = f"#'{sheet.ws.title}'!A1"
            link.font = Font(name=FONT, size=10, bold=True, color=NAVY_3, underline=None)
            link.fill = _fill(SURFACE_4)
            link.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            link.border = Border(bottom=_side(), right=_side())

            ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
            about = ws.cell(row=r, column=3, value=sheet.description or sheet.subtitle or sheet.title)
            about.font = Font(name=FONT, size=9, color=TEXT_2)
            about.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
            for col in range(3, 6):
                cell = ws.cell(row=r, column=col)
                cell.border = Border(bottom=_side())
                if i % 2:
                    cell.fill = _fill(ZEBRA_ROW)
            ws.row_dimensions[r].height = 28
            r += 1

        foot = ws.cell(row=r + 1, column=2,
                       value=f"Файл сформирован автоматически · {datetime.now():%d.%m.%Y %H:%M}")
        foot.font = Font(name=FONT, size=8, color=MUTED)

    def save(self) -> bytes:
        self.wb.remove(self._scratch)
        self._toc()
        for ws in self.wb.worksheets:
            for view in ws.views.sheetView:
                view.tabSelected = False
        self.wb.active = 0
        self.wb.worksheets[0].views.sheetView[0].tabSelected = True

        out = BytesIO()
        self.wb.save(out)
        return out.getvalue()


def period_text(start: date | None, end: date | None) -> str:
    if not start or not end:
        return ""
    return f"период: {start:%d.%m.%Y} — {end:%d.%m.%Y}"
