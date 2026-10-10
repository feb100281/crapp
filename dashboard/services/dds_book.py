"""
Листы отчёта ДДС в Excel после сводной таблицы:

    Расшифровка            статья / подстатья / контрагент — отдельными колонками
    Контрагенты            по одному на строку: поступило, оплатили, сальдо, статьи
    Карточка контрагента   выбор из списка → помесячно и по статьям (формулы по «Операциям»)
    Операции               исходные строки; контрагент приведён к одному названию
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date

from django.db.models import Count, Max, Q, Sum
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.reports.xlsx import (FMT_CODE, FMT_MONEY, FONT, HEADER_ROW, MUTED, NAVY, NUMERIC, SURFACE_3, TEXT,
                               ZEBRA_ROW, Col, Row)

from .cp_report import SKIP, clean_name, key_of

MONTHS = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август",
          "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
NO_CF = "Не разнесено"
INPUT_FILL = "FFF4D6"
SOURCE_NAMES = {"BANK": "Разнесено", "UNALLOC": "Не разнесено", "FEE": "Комиссия", "DEBT": "Удержание",
                "MANUAL": "Проводка", "FX": "Курсовая"}

# колонки листа «Операции» — на них ссылаются формулы карточки
OPS_COLS = [
    ("date", Col("Дата", kind="date", width=12)),
    ("account_name", Col("Счёт", width=28)),
    ("bank_name", Col("Банк", width=24)),
    ("currency", Col("Валюта", width=8)),
    ("source", Col("Тип", width=14)),
    ("activity_name", Col("Деятельность", width=18)),
    ("direction_name", Col("Направление", width=13)),
    ("cf_code", Col("Код", kind="code", width=9)),
    ("article_name", Col("Статья", width=30)),
    ("cf_name", Col("Подстатья", width=30)),
    ("cp", Col("Контрагент", width=38)),
    ("inn", Col("ИНН", kind="code", width=13)),
    ("cp_name", Col("Контрагент в выписке", width=38)),
    ("description", Col("Назначение платежа", width=70)),
    ("amount_cur", Col("Сумма, вал.", kind="money_dec", width=16)),
    ("rate", Col("Курс", kind="money_dec", width=10)),
    ("amount_rub", Col("Сумма, ₽", kind="money_dec", width=17, total=True)),
]


def _letter(field: str) -> str:
    return get_column_letter([f for f, _ in OPS_COLS].index(field) + 1)


def _names(flows) -> dict[str, str]:
    """Ключ контрагента (ИНН или название) → одно название: самое частое из выписок."""
    seen: dict[str, Counter] = defaultdict(Counter)
    for r in flows.exclude(cp_name__isnull=True).values("inn", "cp_name").annotate(n=Count("id")).order_by():
        seen[key_of(r["inn"], r["cp_name"])][clean_name(r["cp_name"])] += r["n"]
    names = {k: c.most_common(1)[0][0] for k, c in seen.items()}
    # одинаковые названия у разных ИНН — различаем ИНН в скобках
    count = Counter(names.values())
    return {k: (f"{n} ({k[4:]})" if count[n] > 1 and k.startswith("inn:") else n) for k, n in names.items()}


def _header_cells(sheet, row: int, start: int, cols: list[Col]) -> None:
    for i, col in enumerate(cols):
        cell = sheet.ws.cell(row=row, column=start + i, value=col.title)
        cell._style = sheet.book.style(bold=True, color="FFFFFF", fill=NAVY, wrap=True, vcenter=True,
                                       align="right" if col.kind in NUMERIC else "left", right=True)
        sheet.ws.column_dimensions[get_column_letter(start + i)].width = col.width or 14


def add_sheets(book, flows, year: int | None) -> None:
    names = _names(flows)
    _detail(book, flows, names)
    cps = _counterparties(book, flows, names)
    _card(book, flows, cps, year)
    _operations(book, flows, names)


# ---------------------------------------------------------------- Расшифровка

def _detail(book, flows, names) -> None:
    groups = defaultdict(list)
    for g in (flows.exclude(source="FX")
              .values("activity", "activity_name", "direction_name", "cf_code", "article_name", "cf_name",
                      "cp_name", "inn")
              .annotate(s=Sum("amount_rub"), n=Count("id")).order_by()):
        groups[(g["activity"] or 9, g["direction_name"] or "", g["cf_code"] or "", g["article_name"] or "",
                g["cf_name"] or NO_CF, g["activity_name"] or "")].append(g)

    rows = []
    for key in sorted(groups, key=lambda k: (k[0], k[1] != "Поступления", k[2])):
        activity, direction, code, article, sub, act_name = key
        if not article or article == sub:
            article, sub = sub, ""
        by_cp: dict[str, list] = defaultdict(lambda: ["", 0, 0.0])
        for g in groups[key]:
            k = key_of(g["inn"], g["cp_name"])
            item = by_cp[names.get(k, "без контрагента") if g["cp_name"] else "без контрагента"]
            item[0] = item[0] or (g["inn"] or "")
            item[1] += g["n"]
            item[2] += g["s"] or 0
        cps = sorted(by_cp.items(), key=lambda kv: -abs(kv[1][2]))
        rows.append(Row([code, act_name, direction, article, sub, f"Итого · контрагентов: {len(cps)}", "",
                         sum(c[1] for _, c in cps), sum(c[2] for _, c in cps)],
                        level="warn" if activity == 9 else 3, collapsed=True))
        rows += [Row([code, act_name, direction, article, sub, cp, inn, n, s], outline=1, hidden=True)
                 for cp, (inn, n, s) in cps]

    book.sheet(
        "Расшифровка", "Расшифровка статей по контрагентам",
        subtitle="Строка статьи — итог; плюсик слева раскрывает контрагентов по убыванию суммы",
        description="Статья, подстатья и контрагент — в отдельных колонках: удобно фильтровать и сортировать",
    ).table(
        [Col("Код", kind="code", width=9), Col("Деятельность", width=20), Col("Направление", width=13),
         Col("Статья", width=34), Col("Подстатья", width=34), Col("Контрагент", width=44),
         Col("ИНН", kind="code", width=13), Col("Операций", kind="int", width=10),
         Col("Сумма, ₽", kind="money", width=17, total=True)],
        rows, freeze_cols=0, autofilter=True,
    )


# ---------------------------------------------------------------- Контрагенты

def _counterparties(book, flows, names) -> list[str]:
    stat: dict[str, dict] = {}
    for r in (flows.exclude(SKIP).exclude(cp_name__isnull=True)
              .values("inn", "cp_name", "article_name", "cf_name")
              .annotate(inflow=Sum("amount_rub", filter=Q(amount_rub__gt=0), default=0),
                        outflow=Sum("amount_rub", filter=Q(amount_rub__lt=0), default=0),
                        n=Count("id"), last=Max("date")).order_by()):
        k = key_of(r["inn"], r["cp_name"])
        s = stat.setdefault(k, {"name": names[k], "inn": r["inn"] or "", "in": 0.0, "out": 0.0, "n": 0,
                                "last": r["last"], "arts": Counter()})
        s["in"] += r["inflow"] or 0
        s["out"] += r["outflow"] or 0
        s["n"] += r["n"]
        s["last"] = max(s["last"], r["last"])
        s["arts"][r["cf_name"] or r["article_name"] or NO_CF] += abs((r["inflow"] or 0) + (r["outflow"] or 0))
    items = sorted(stat.values(), key=lambda s: -(s["in"] - s["out"]))

    book.sheet(
        "Контрагенты", "Контрагенты за период",
        subtitle="Кто сколько нам заплатил и сколько мы ему. Фильтр в шапке — найти контрагента; "
                 "подробно — на листе «Карточка контрагента»",
        description="Один контрагент — одна строка: поступило, оплатили, сальдо, основные статьи",
    ).table(
        [Col("Контрагент", width=46), Col("ИНН", kind="code", width=13),
         Col("Поступило, ₽", kind="money", width=17), Col("Оплатили, ₽", kind="money", width=17),
         Col("Сальдо, ₽", kind="money", width=17, total=True), Col("Операций", kind="int", width=10),
         Col("Последняя", kind="date", width=12), Col("Основные статьи", width=60)],
        [Row([s["name"], s["inn"], s["in"], s["out"], s["in"] + s["out"], s["n"], s["last"],
              " · ".join(a for a, _ in s["arts"].most_common(3))]) for s in items],
        total=Row(["ИТОГО", "", sum(s["in"] for s in items), sum(s["out"] for s in items),
                   sum(s["in"] + s["out"] for s in items), sum(s["n"] for s in items), None, ""]),
        freeze_cols=1, autofilter=True,
    )
    return [s["name"] for s in items]


# ---------------------------------------------------------------- Карточка

def _card(book, flows, cps: list[str], year: int | None) -> None:
    if not cps:
        return
    sheet = book.sheet(
        "Карточка контрагента", "Карточка контрагента",
        description="Выберите контрагента из списка — помесячно и по статьям: поступило, оплатили, сальдо",
    )
    ws = sheet.ws
    ops, last_cp = "'Операции'", 7 + len(cps) - 1
    amt, cp, d, code = (f"{ops}!${_letter(f)}:${_letter(f)}" for f in ("amount_rub", "cp", "date", "cf_code"))
    sel = "$B$4"

    # помесячно: A–E
    if year:
        periods = [(MONTHS[m - 1], date(year, m, 1), date(year + (m == 12), m % 12 + 1, 1)) for m in range(1, 13)]
    else:
        ys = sorted({y for y in flows.order_by().values_list("year", flat=True).distinct()})
        periods = [(f"{y} год", date(y, 1, 1), date(y + 1, 1, 1)) for y in ys]

    def sumifs(extra: str = "", sign: str = "") -> str:
        cond = f',{amt},"{sign}0"' if sign else ""
        return f"SUMIFS({amt},{cp},{sel}{extra}{cond})"

    rows = []
    for label, d1, d2 in periods:
        span = f',{d},">="&DATE({d1.year},{d1.month},1),{d},"<"&DATE({d2.year},{d2.month},1)'
        r = HEADER_ROW + 1 + len(rows)
        rows.append(Row([label, "=" + sumifs(span, ">"), "=" + sumifs(span, "<"), f"=B{r}+C{r}",
                         f"=COUNTIFS({cp},{sel}{span})"]))
    first, last = HEADER_ROW + 1, HEADER_ROW + len(rows)
    sheet.table(
        [Col("Период", width=14), Col("Поступило, ₽", kind="money", width=17),
         Col("Оплатили, ₽", kind="money", width=17), Col("Сальдо, ₽", kind="money", width=17, total=True),
         Col("Операций", kind="int", width=10)],
        rows,
        total=Row(["ИТОГО", f"=SUM(B{first}:B{last})", f"=SUM(C{first}:C{last})", f"=SUM(D{first}:D{last})",
                   f"=SUM(E{first}:E{last})"]),
        freeze_cols=0,
    )

    # выбор контрагента — в строке 4 вместо параметров
    ws.unmerge_cells(start_row=4, start_column=1, end_row=4, end_column=5)
    ws.merge_cells("B4:E4")
    ws["A4"] = "Контрагент ▸"
    ws["A4"].font = Font(name=FONT, size=10, bold=True, color=NAVY)
    ws["B4"] = cps[0]
    ws["B4"].font = Font(name=FONT, size=11, bold=True)
    for c in "BCDE":
        ws[f"{c}4"].fill = PatternFill("solid", start_color=INPUT_FILL, end_color=INPUT_FILL)
    ws.row_dimensions[4].height = 22
    ws["A3"] = (f'=IFERROR("ИНН "&INDEX(\'Контрагенты\'!$B$7:$B${last_cp},'
                f'MATCH({sel},\'Контрагенты\'!$A$7:$A${last_cp},0))&"  ·  ","")'
                f'&"щёлкните жёлтую ячейку и выберите контрагента из списка"')
    dv = DataValidation(type="list", formula1=f"'Контрагенты'!$A$7:$A${last_cp}", allow_blank=False,
                        showErrorMessage=True, errorTitle="Контрагент",
                        error="Выберите контрагента из списка")
    dv.add("B4")
    ws.add_data_validation(dv)

    # по статьям: G–L, та же шапка
    arts = list(flows.exclude(SKIP).exclude(cp_name__isnull=True)
                .values("cf_code", "article_name", "cf_name").annotate(n=Count("id"))
                .order_by("cf_code"))
    cols = [Col("Код", kind="code", width=9), Col("Статья", width=32), Col("Подстатья", width=32),
            Col("Поступило, ₽", kind="money", width=17), Col("Оплатили, ₽", kind="money", width=17),
            Col("Сальдо, ₽", kind="money", width=17, total=True)]
    start = 7
    _header_cells(sheet, HEADER_ROW, start, cols)
    ws.column_dimensions[get_column_letter(start - 1)].width = 3
    for i, a in enumerate(arts):
        r = HEADER_ROW + 1 + i
        article, sub = a["article_name"] or "", a["cf_name"] or NO_CF
        if not article or article == sub:
            article, sub = sub, ""
        crit = f',{code},"{a["cf_code"]}"' if a["cf_code"] else f',{code},""'
        p, q = get_column_letter(start + 3), get_column_letter(start + 4)
        values = [int(a["cf_code"]) if (a["cf_code"] or "").isdigit() else (a["cf_code"] or ""), article, sub,
                  "=" + sumifs(crit, ">"), "=" + sumifs(crit, "<"), f"={p}{r}+{q}{r}"]
        for j, (col, value) in enumerate(zip(cols, values)):
            cell = ws.cell(row=r, column=start + j, value=value)
            numeric = col.kind in NUMERIC
            cell._style = sheet.book.style(
                color=MUTED if col.kind == "code" else TEXT, bold=col.total,
                fill=SURFACE_3 if col.total else (ZEBRA_ROW if r % 2 == 0 else None),
                fmt=FMT_MONEY if numeric else (FMT_CODE if col.kind == "code" else None),
                align="right" if numeric else "left", right=True)
    if arts:
        lo, hi = HEADER_ROW + 1, HEADER_ROW + len(arts)
        p, q = get_column_letter(start + 3), get_column_letter(start + 4)
        ws.conditional_formatting.add(
            f"{get_column_letter(start)}{lo}:{get_column_letter(start + 5)}{hi}",
            FormulaRule(formula=[f"AND(${p}{lo}=0,${q}{lo}=0)"], font=Font(color="C8C8C8")))


# ---------------------------------------------------------------- Операции

def _operations(book, flows, names) -> None:
    fields = [f for f, _ in OPS_COLS if f not in ("cp",)]

    def rows():
        for o in flows.order_by("date", "id").values(*fields).iterator():
            k = key_of(o["inn"], o["cp_name"])
            o["cp"] = names.get(k, "") if o["cp_name"] else ""
            o["source"] = SOURCE_NAMES.get(o["source"], o["source"])
            o["description"] = (o["description"] or "")[:500]
            yield Row([o[f] if o[f] is not None else ("" if f not in ("amount_cur", "rate", "amount_rub") else None)
                       for f, _ in OPS_COLS])

    book.sheet(
        "Операции", "Операции за период",
        subtitle="Исходные строки отчёта. «Контрагент» — одно название на ИНН, «в выписке» — как написал банк",
        description="Плоская таблица всех операций периода — для фильтров и своих сводных",
    ).table([c for _, c in OPS_COLS], rows(), autofilter=True)

